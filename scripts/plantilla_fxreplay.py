"""Plantilla de FX Replay para validar Momento 09:40 (stop 0,40%, 1:2) a mano: 60 sesiones de 2022."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from src import loader

b = loader.cargar_cfd([loader.RAIZ / "data/nq_cfd_2022.csv.gz"], verbose=False)
dias = [t.date() for t in b[b.hhmm == 940].index if t.year == 2022][:60]

wb = Workbook()
ins = wb.active; ins.title = "Instrucciones"
texto = [
    "VALIDACIÓN A MANO · MOMENTO 09:40 · NQ / US100 · 60 sesiones seguidas de 2022",
    "",
    "Antes de empezar",
    "1. Gráfico en HORA DE NUEVA YORK (no la de España ni UTC). Velas de 1 minuto.",
    "2. Haz las 60 sesiones SEGUIDAS, sin saltarte ninguna, aunque el día 'no te guste'.",
    "3. No mires el día siguiente antes de apuntar el resultado del actual.",
    "",
    "Cada día",
    "A. Apunta el CIERRE de la vela de las 09:09 (columna B).",
    "B. Apunta el CIERRE de la vela de las 09:39 (columna C).",
    "   · La hoja te dice la dirección: 09:39 > 09:09 → LARGO · 09:39 < 09:09 → CORTO.",
    "C. Entrada a mercado en la APERTURA de la vela de las 09:40 (columna D).",
    "   · La hoja calcula el stop (0,40% del precio) y el TP (el doble), columnas F-H.",
    "D. Avanza el replay hasta que pase una de tres cosas y elige en la columna I:",
    "   · TP  = toca el objetivo · SL = toca el stop · CIERRE = no toca nada antes de las 15:59.",
    "   · Si en la MISMA vela de 1 minuto se tocan TP y SL, apunta SL.",
    "E. Si es CIERRE, apunta el precio de cierre de la vela de las 15:59 en la columna J.",
    "",
    "La hoja calcula la R de cada día (TP = +2, SL = −1, CIERRE = lo que corresponda).",
    "La comisión se descuenta aparte en la hoja Resumen.",
    "",
    "Qué esperar (sacado del backtest): unas +5R en 60 días, pero lo normal va de −9R a +19R.",
    "Un resultado dentro de ese rango es compatible con el backtest. Si sale MUY por debajo,",
    "puede haber un error en el código: avísame.",
    "Cuando acabes, pásame el Excel y lo cruzo día a día con el backtest.",
]
for i, t in enumerate(texto, 1):
    ins.cell(i, 1, t)
ins["A1"].font = Font(bold=True, size=13)
for r in (3, 8):
    ins.cell(r, 1).font = Font(bold=True)
ins.column_dimensions["A"].width = 100

ws = wb.create_sheet("Operaciones")
cab = ["Fecha", "Cierre 09:09", "Cierre 09:39", "Apertura 09:40 (entrada)", "Dirección", "Distancia stop (0,40%)",
       "Precio STOP", "Precio TP (2R)", "Resultado (TP/SL/CIERRE)", "Cierre 15:59 (si CIERRE)", "R", "Notas"]
for j, c in enumerate(cab, 1):
    x = ws.cell(1, j, c); x.font = Font(bold=True, color="FFFFFF"); x.fill = PatternFill("solid", fgColor="305496")
    x.alignment = Alignment(wrap_text=True, vertical="center")
dv = DataValidation(type="list", formula1='"TP,SL,CIERRE"', allow_blank=True); ws.add_data_validation(dv)
for i, d in enumerate(dias, 2):
    ws.cell(i, 1, d).number_format = "DD/MM/YYYY"
    ws.cell(i, 5, f'=IF(OR(B{i}="",C{i}=""),"",IF(C{i}>B{i},"LARGO",IF(C{i}<B{i},"CORTO","NADA")))')
    ws.cell(i, 6, f'=IF(D{i}="","",ROUND(D{i}*0.004,2))')
    ws.cell(i, 7, f'=IF(OR(D{i}="",E{i}="NADA"),"",IF(E{i}="LARGO",D{i}-F{i},D{i}+F{i}))')
    ws.cell(i, 8, f'=IF(OR(D{i}="",E{i}="NADA"),"",IF(E{i}="LARGO",D{i}+2*F{i},D{i}-2*F{i}))')
    dv.add(ws.cell(i, 9))
    ws.cell(i, 11, f'=IF(I{i}="TP",2,IF(I{i}="SL",-1,IF(AND(I{i}="CIERRE",J{i}<>""),ROUND(IF(E{i}="LARGO",J{i}-D{i},D{i}-J{i})/F{i},3),"")))')
for j, w in enumerate([12, 12, 12, 14, 10, 12, 12, 12, 14, 14, 8, 30], 1):
    ws.column_dimensions[chr(64 + j)].width = w
ws.freeze_panes = "A2"

rs = wb.create_sheet("Resumen")
n = len(dias) + 1
filas = [
    ("Días apuntados", f'=COUNTA(Operaciones!I2:I{n})'),
    ("TP", f'=COUNTIF(Operaciones!I2:I{n},"TP")'),
    ("SL", f'=COUNTIF(Operaciones!I2:I{n},"SL")'),
    ("Cierres 15:59", f'=COUNTIF(Operaciones!I2:I{n},"CIERRE")'),
    ("R total (sin comisión)", f'=SUM(Operaciones!K2:K{n})'),
    ("R media por operación", '=IF(B1=0,"",B5/B1)'),
    ("Comisión aprox. (0,87 pts / stop medio)", f'=IF(B1=0,"",B1*0.87/AVERAGE(Operaciones!F2:F{n}))'),
    ("R total con comisión", '=IF(B1=0,"",B5-B7)'),
    ("Esperado del backtest", "≈ +5R en 60 días (rango normal −9R a +19R)"),
]
for i, (k, f) in enumerate(filas, 1):
    rs.cell(i, 1, k).font = Font(bold=True); rs.cell(i, 2, f)
rs.column_dimensions["A"].width = 42; rs.column_dimensions["B"].width = 44
wb.save(loader.RAIZ / "fxreplay/plantilla_momento_0940_2022.xlsx")
print("ok", len(dias), dias[0], dias[-1])
